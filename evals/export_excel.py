from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook



def load_json(path: Path) -> list[dict]:
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))



def export_results_to_excel(results_dir: Path, output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_sheet = workbook.create_sheet("summaries")
    summary_headers = [
        "provider",
        "model",
        "total_queries",
        "avg_latency_ms",
        "estimated_total_input_cost",
        "estimated_total_output_cost",
        "estimated_total_cost",
        "avg_prompt_tokens",
        "avg_completion_tokens",
        "avg_total_tokens",
        "avg_score",
    ]
    summary_sheet.append(summary_headers)

    result_sheet = workbook.create_sheet("results")
    result_headers = [
        "query_id",
        "provider",
        "model",
        "prompt",
        "response_text",
        "latency_ms",
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
        "estimated_input_cost",
        "estimated_output_cost",
        "total_cost",
        "parameter_size",
        "category",
        "language",
        "title",
        "mode",
        "expected_intent",
        "expected_data_source",
    ]
    result_sheet.append(result_headers)

    score_sheet = workbook.create_sheet("scores")
    score_headers = ["query_id", "provider", "model", "category", "score", "max_score", "rationale"]
    score_sheet.append(score_headers)

    judge_sheet = workbook.create_sheet("judge_scores")
    judge_headers = [
        "query_id",
        "provider",
        "model",
        "category",
        "judge_provider",
        "judge_model",
        "judge_version",
        "rubric_version",
        "overall_score",
        "decision",
        "confidence",
        "judge_latency_ms",
        "judge_total_tokens",
        "judge_total_cost",
    ]
    judge_sheet.append(judge_headers)

    e2e_case_sheet = workbook.create_sheet("e2e_cases")
    e2e_case_headers = [
        "run_id",
        "case_id",
        "query_id",
        "category",
        "language",
        "title",
        "mode",
        "operation",
        "query_zh",
        "expected_intent",
        "expected_data_source",
        "profile_id",
        "fixture_mode",
        "provider",
        "model",
        "request_text",
        "response_text",
        "http_status",
        "api_status",
        "latency_ms",
        "estimated_total_cost",
        "score",
        "max_score",
        "decision",
        "degraded",
        "degraded_reason",
        "retrieval_strategy",
        "rag_backend_ready",
        "rag_query_status",
        "supported",
        "notes",
    ]
    e2e_case_sheet.append(e2e_case_headers)

    e2e_summary_sheet = workbook.create_sheet("e2e_summary")
    e2e_summary_headers = [
        "model",
        "case_count",
        "avg_score",
        "pass_count",
        "pass_rate",
        "avg_latency_ms",
        "categories",
        "operations",
    ]
    e2e_summary_sheet.append(e2e_summary_headers)

    routing_case_sheet = workbook.create_sheet("routing_cases")
    routing_case_headers = [
        "candidate_name",
        "case_id",
        "bucket",
        "language",
        "query",
        "query_zh",
        "expected_intent",
        "rule_intent",
        "final_intent",
        "intent_correct",
        "domain_supported",
        "hard_deny",
        "needs_fallback",
        "fallback_used",
        "fallback_reason",
        "boundary_topic",
        "out_of_scope_subtype",
        "latency_ms",
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
        "estimated_total_cost",
        "fallback_provider",
        "fallback_model",
        "notes",
    ]
    routing_case_sheet.append(routing_case_headers)

    routing_summary_sheet = workbook.create_sheet("routing_summary")
    routing_summary_headers = [
        "dataset_name",
        "dataset_version",
        "candidate_name",
        "fallback_enabled",
        "total_cases",
        "correct_cases",
        "accuracy",
        "fallback_needed_cases",
        "fallback_used_cases",
        "fallback_hit_rate",
        "needs_fallback_accuracy",
        "boundary_accuracy",
        "manual_route_edit_accuracy",
        "conflict_accuracy",
        "avg_latency_ms",
        "estimated_total_cost",
    ]
    routing_summary_sheet.append(routing_summary_headers)

    translation_case_sheet = workbook.create_sheet("translation_cases")
    translation_case_headers = [
        "run_id",
        "suite",
        "dataset_name",
        "dataset_version",
        "case_id",
        "question_type",
        "category",
        "direction",
        "source_language",
        "target_language",
        "user_language",
        "provider",
        "model",
        "query",
        "response_text",
        "status",
        "reason",
        "execution_path",
        "latency_ms",
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
        "estimated_input_cost",
        "estimated_output_cost",
        "estimated_total_cost",
        "downstream_intent",
        "downstream_confidence",
        "score",
        "max_score",
        "passed_checks",
        "failed_checks",
    ]
    translation_case_sheet.append(translation_case_headers)

    translation_compare_sheet = workbook.create_sheet("translation_model_compare")
    translation_compare_headers = [
        "run_id",
        "suite",
        "language",
        "question_type",
        "direction",
        "provider",
        "model",
        "case_count",
        "avg_score",
        "avg_latency_ms",
        "sum_total_tokens",
        "estimated_total_cost",
        "degraded_cases",
    ]
    translation_compare_sheet.append(translation_compare_headers)

    for summary_file in sorted(results_dir.glob("*_summary.json")):
        summary = json.loads(summary_file.read_text(encoding="utf-8"))
        if isinstance(summary, dict):
            summary_sheet.append([summary.get(header) for header in summary_headers])

    for result_file in sorted(results_dir.glob("*_results.json")):
        for item in load_json(result_file):
            metadata = item.get("metadata", {})
            result_sheet.append(
                [
                    item.get("query_id"),
                    item.get("provider"),
                    item.get("model"),
                    item.get("prompt"),
                    item.get("response_text"),
                    item.get("latency_ms"),
                    item.get("prompt_tokens"),
                    item.get("completion_tokens"),
                    item.get("total_tokens"),
                    item.get("estimated_input_cost"),
                    item.get("estimated_output_cost"),
                    item.get("total_cost"),
                    item.get("parameter_size"),
                    metadata.get("category"),
                    metadata.get("language"),
                    metadata.get("title"),
                    metadata.get("mode"),
                    metadata.get("expected_intent"),
                    metadata.get("expected_data_source"),
                ]
            )

    for score_file in sorted(results_dir.glob("*_scores.json")):
        for item in load_json(score_file):
            score_sheet.append([item.get(header) for header in score_headers])

    for judge_file in sorted(results_dir.glob("*_judge_scores.json")):
        for item in load_json(judge_file):
            judge_sheet.append([item.get(header) for header in judge_headers])

    e2e_rows: list[dict] = []
    for case_file in sorted(results_dir.glob("e2e_qa_benchmark_run*_cases.json")):
        for item in load_json(case_file):
            e2e_rows.append(item)
            e2e_case_sheet.append([item.get(header) for header in e2e_case_headers])

    model_stats: dict[str, dict] = {}
    for item in e2e_rows:
        model = item.get("model") or "unknown"
        stats = model_stats.setdefault(
            model,
            {
                "case_count": 0,
                "score_sum": 0.0,
                "score_count": 0,
                "pass_count": 0,
                "latency_sum": 0.0,
                "latency_count": 0,
                "categories": set(),
                "operations": set(),
            },
        )
        stats["case_count"] += 1
        score = item.get("score")
        if isinstance(score, (int, float)):
            stats["score_sum"] += float(score)
            stats["score_count"] += 1
        if item.get("decision") == "pass":
            stats["pass_count"] += 1
        latency = item.get("latency_ms")
        if isinstance(latency, (int, float)):
            stats["latency_sum"] += float(latency)
            stats["latency_count"] += 1
        category = item.get("category")
        if category:
            stats["categories"].add(str(category))
        operation = item.get("operation")
        if operation:
            stats["operations"].add(str(operation))

    for model, stats in sorted(model_stats.items()):
        case_count = stats["case_count"]
        avg_score = stats["score_sum"] / stats["score_count"] if stats["score_count"] else None
        avg_latency = stats["latency_sum"] / stats["latency_count"] if stats["latency_count"] else None
        pass_rate = stats["pass_count"] / case_count if case_count else None
        e2e_summary_sheet.append(
            [
                model,
                case_count,
                avg_score,
                stats["pass_count"],
                pass_rate,
                avg_latency,
                ", ".join(sorted(stats["categories"])),
                ", ".join(sorted(stats["operations"])),
            ]
        )

    for case_file in sorted(results_dir.glob("routing_*_cases.json")):
        for item in load_json(case_file):
            routing_case_sheet.append([item.get(header) for header in routing_case_headers])

    for summary_file in sorted(results_dir.glob("routing_*_summary.json")):
        if summary_file.name == "routing_sweep_summary.json":
            continue
        summary = json.loads(summary_file.read_text(encoding="utf-8"))
        if not isinstance(summary, dict) or "candidate_name" not in summary:
            continue
        routing_summary_sheet.append([summary.get(header) for header in routing_summary_headers])

    translation_rows: list[dict] = []
    for result_file in sorted(results_dir.glob("translation_*_results.json")):
        for item in load_json(result_file):
            translation_rows.append(item)
            translation_case_sheet.append(
                [
                    item.get("run_id"),
                    item.get("suite"),
                    item.get("dataset_name"),
                    item.get("dataset_version"),
                    item.get("case_id") or item.get("id"),
                    item.get("question_type"),
                    item.get("category"),
                    item.get("direction"),
                    item.get("source_language"),
                    item.get("target_language"),
                    item.get("user_language"),
                    item.get("provider"),
                    item.get("model"),
                    item.get("query"),
                    item.get("response_text"),
                    item.get("status"),
                    item.get("reason"),
                    item.get("execution_path"),
                    item.get("latency_ms"),
                    item.get("prompt_tokens"),
                    item.get("completion_tokens"),
                    item.get("total_tokens"),
                    item.get("estimated_input_cost"),
                    item.get("estimated_output_cost"),
                    item.get("estimated_total_cost"),
                    item.get("downstream_intent"),
                    item.get("downstream_confidence"),
                    item.get("score"),
                    item.get("max_score"),
                    ", ".join(item.get("passed_checks") or []),
                    ", ".join(item.get("failed_checks") or []),
                ]
            )

    translation_stats: dict[tuple[str, str, str, str, str, str, str], dict] = {}
    for item in translation_rows:
        key = (
            str(item.get("run_id") or "unknown"),
            str(item.get("suite") or "unknown"),
            str(item.get("target_language") or item.get("source_language") or "unknown"),
            str(item.get("question_type") or "unknown"),
            str(item.get("direction") or "unknown"),
            str(item.get("provider") or "unknown"),
            str(item.get("model") or "unknown"),
        )
        stats = translation_stats.setdefault(
            key,
            {
                "case_count": 0,
                "score_sum": 0.0,
                "latency_sum": 0.0,
                "sum_total_tokens": 0,
                "estimated_total_cost": 0.0,
                "degraded_cases": 0,
            },
        )
        stats["case_count"] += 1
        stats["score_sum"] += float(item.get("score") or 0.0)
        stats["latency_sum"] += float(item.get("latency_ms") or 0.0)
        stats["sum_total_tokens"] += int(item.get("total_tokens") or 0)
        stats["estimated_total_cost"] += float(item.get("estimated_total_cost") or 0.0)
        if item.get("status") != "ok":
            stats["degraded_cases"] += 1

    for key, stats in sorted(translation_stats.items()):
        run_id, suite, language, question_type, direction, provider, model = key
        case_count = stats["case_count"]
        translation_compare_sheet.append(
            [
                run_id,
                suite,
                language,
                question_type,
                direction,
                provider,
                model,
                case_count,
                stats["score_sum"] / case_count if case_count else None,
                stats["latency_sum"] / case_count if case_count else None,
                stats["sum_total_tokens"],
                stats["estimated_total_cost"],
                stats["degraded_cases"],
            ]
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export benchmark JSON artifacts to xlsx")
    parser.add_argument("--results-dir")
    parser.add_argument("--output")
    args = parser.parse_args()

    project_root = Path(__file__).resolve().parent.parent
    default_results_dir = project_root / "evals" / "results"
    results_dir = Path(args.results_dir) if args.results_dir else default_results_dir
    output_path = Path(args.output) if args.output else (results_dir / "benchmark_results.xlsx")
    export_results_to_excel(results_dir, output_path)

from __future__ import annotations

import argparse
import asyncio
import json
from collections import defaultdict
from pathlib import Path
from statistics import mean, median
from time import perf_counter
from typing import Any

from openpyxl import Workbook

from yoyo.core.config import get_settings
from yoyo.modules.knowledge.hybrid_context_builder import _resolve_navigation_slots
from yoyo.modules.knowledge.navigation_retriever import get_navigation_context
from yoyo.modules.knowledge.schemas import NavigationSlotPayload
from yoyo.modules.qa.formatters import format_navigation_text_answer
from yoyo.modules.qa.intent_router import score_intent


async def _resolve_case(case: dict[str, Any], *, check_amap: bool = False) -> dict[str, Any]:
    case_started_at = perf_counter()
    query = str(case.get("query") or "")
    session_context = dict(case.get("session_context") or {})
    expected_intent = case.get("expected_intent")
    intent_started_at = perf_counter()
    router_result = score_intent(query, raw_query=query)
    intent_latency_ms = _elapsed_ms(intent_started_at)
    actual_intent = router_result.get("intent")
    row: dict[str, Any] = {
        "case_id": case.get("case_id"),
        "category": case.get("category"),
        "sub_category": case.get("sub_category"),
        "query": query,
        "language": case.get("language"),
        "mode_case": case.get("mode_case"),
        "distance_bucket": case.get("distance_bucket"),
        "expected_intent": expected_intent,
        "actual_intent": actual_intent,
        "intent_correct": actual_intent == expected_intent,
        "intent_latency_ms": intent_latency_ms,
        "notes": [],
    }

    if expected_intent != "navigation_text":
        row.update(
            {
                "expected_request_kind": None,
                "actual_request_kind": None,
                "expected_origin": None,
                "actual_origin": None,
                "expected_destinations": [],
                "actual_destinations": [],
                "expected_source": None,
                "actual_source": None,
                "slot_correct": actual_intent == expected_intent,
                "amap_checked": False,
                "amap_status": None,
                "amap_final_answer": None,
                "amap_requested_mode": None,
                "amap_final_mode": None,
                "amap_mode_fallback_used": None,
                "amap_mode_fallback_reason": None,
                "amap_leg_final_modes": [],
                "amap_leg_mode_fallbacks": [],
                "amap_result_reasonable": None,
                "amap_latency_ms": None,
                "total_latency_ms": _elapsed_ms(case_started_at),
            }
        )
        row["score"] = 1.0 if row["intent_correct"] else 0.0
        row["decision"] = "pass" if row["intent_correct"] else "fail"
        return row

    slot_started_at = perf_counter()
    slot_result, slot_debug = await _resolve_navigation_slots(query, attraction=None, session_context=session_context)
    slot_latency_ms = _elapsed_ms(slot_started_at)
    actual_destinations = list(slot_result.destinations)
    expected_destinations = list(case.get("expected_destinations") or [])
    actual_origin_place_id = slot_result.origin_place.place_id if slot_result.origin_place else None
    actual_destination_place_ids = [place.place_id for place in slot_result.destination_places if place.place_id]
    expected_origin_place_id = case.get("expected_origin_place_id")
    expected_destination_place_ids = list(case.get("expected_destination_place_ids") or [])
    request_kind_correct = slot_result.request_kind == case.get("expected_request_kind")
    origin_correct = slot_result.origin == case.get("expected_origin")
    destinations_correct = actual_destinations == expected_destinations
    source_correct = slot_result.source == case.get("expected_source")
    expected_mode = case.get("expected_mode")
    expected_mode_source = case.get("expected_mode_source")
    expected_transit_preference = case.get("expected_transit_preference")
    mode_correct = expected_mode is None or slot_result.mode == expected_mode
    mode_source_correct = expected_mode_source is None or slot_result.mode_source == expected_mode_source
    transit_preference_correct = (
        expected_transit_preference is None or slot_result.transit_preference == expected_transit_preference
    )
    origin_place_correct = expected_origin_place_id is None or actual_origin_place_id == expected_origin_place_id
    destination_places_correct = (
        not expected_destination_place_ids or actual_destination_place_ids == expected_destination_place_ids
    )
    slot_correct = all(
        [
            request_kind_correct,
            origin_correct,
            destinations_correct,
            source_correct,
            mode_correct,
            mode_source_correct,
            transit_preference_correct,
            origin_place_correct,
            destination_places_correct,
        ]
    )

    row.update(
        {
            "expected_request_kind": case.get("expected_request_kind"),
            "actual_request_kind": slot_result.request_kind,
            "expected_origin": case.get("expected_origin"),
            "actual_origin": slot_result.origin,
            "expected_destinations": expected_destinations,
            "actual_destinations": actual_destinations,
            "expected_origin_place_id": expected_origin_place_id,
            "actual_origin_place_id": actual_origin_place_id,
            "expected_destination_place_ids": expected_destination_place_ids,
            "actual_destination_place_ids": actual_destination_place_ids,
            "origin_place_source": slot_result.origin_place.source if slot_result.origin_place else None,
            "destination_place_sources": [place.source for place in slot_result.destination_places],
            "expected_source": case.get("expected_source"),
            "actual_source": slot_result.source,
            "expected_mode": expected_mode,
            "actual_mode": slot_result.mode,
            "expected_mode_source": expected_mode_source,
            "actual_mode_source": slot_result.mode_source,
            "expected_transit_preference": expected_transit_preference,
            "actual_transit_preference": slot_result.transit_preference,
            "slot_debug": slot_debug,
            "slot_latency_ms": slot_latency_ms,
            "slot_correct": slot_correct,
        }
    )

    amap_checked = False
    amap_status = None
    amap_result_reasonable = None
    amap_latency_ms = None
    amap_distance_meters = None
    amap_duration_seconds = None
    amap_leg_count = None
    amap_available_leg_count = None
    amap_step_count = None
    amap_first_steps: list[str] = []
    amap_reason = None
    amap_final_answer = None
    amap_requested_mode = None
    amap_final_mode = None
    amap_mode_fallback_used = None
    amap_mode_fallback_reason = None
    amap_requested_transit_preference = None
    amap_final_transit_vehicle_types: list[str] = []
    amap_transit_preference_satisfied = None
    amap_leg_final_modes: list[str | None] = []
    amap_leg_transit_vehicle_types: list[list[str]] = []
    amap_leg_mode_fallbacks: list[bool] = []
    if check_amap and get_settings().map_api_key and row["intent_correct"] and slot_correct:
        amap_checked = True
        amap_started_at = perf_counter()
        navigation = await get_navigation_context(
            slot_result=NavigationSlotPayload(**slot_result.model_dump())
        )
        amap_latency_ms = _elapsed_ms(amap_started_at)
        amap_status = navigation.status
        amap_distance_meters = navigation.distance_meters
        amap_duration_seconds = navigation.duration_seconds
        amap_leg_count = len(navigation.legs)
        amap_available_leg_count = sum(1 for leg in navigation.legs if leg.status == "available")
        amap_step_count = len(navigation.steps)
        amap_first_steps = [step.instruction for step in navigation.steps[:3]]
        amap_reason = navigation.reason
        amap_requested_mode = navigation.requested_mode
        amap_final_mode = navigation.final_mode
        amap_requested_transit_preference = navigation.requested_transit_preference
        amap_final_transit_vehicle_types = navigation.final_transit_vehicle_types
        amap_transit_preference_satisfied = _transit_preference_satisfied(
            navigation.requested_transit_preference,
            navigation.final_mode,
            navigation.final_transit_vehicle_types,
            navigation.mode_fallback_used,
        )
        amap_mode_fallback_used = navigation.mode_fallback_used
        amap_mode_fallback_reason = navigation.mode_fallback_reason
        amap_leg_final_modes = [leg.final_mode for leg in navigation.legs]
        amap_leg_transit_vehicle_types = [leg.final_transit_vehicle_types for leg in navigation.legs]
        amap_leg_mode_fallbacks = [leg.mode_fallback_used for leg in navigation.legs]
        amap_final_answer = format_navigation_text_answer(
            navigation.origin_name,
            navigation.destination_name,
            navigation.steps,
            navigation.distance_meters,
            navigation.duration_seconds,
            degraded_reason=navigation.reason if navigation.status != "available" else None,
            legs=navigation.legs,
        )
        amap_result_reasonable = bool(
            navigation.status == "available"
            and (navigation.distance_meters or 0) > 0
            and (navigation.duration_seconds or 0) > 0
            and len(navigation.legs) == len(actual_destinations)
            and all(leg.status == "available" for leg in navigation.legs)
            and len(navigation.steps) > 0
            and amap_transit_preference_satisfied is not False
        )
    row.update(
        {
            "amap_checked": amap_checked,
            "amap_status": amap_status,
            "amap_reason": amap_reason,
            "amap_final_answer": amap_final_answer,
            "amap_requested_mode": amap_requested_mode,
            "amap_final_mode": amap_final_mode,
            "amap_requested_transit_preference": amap_requested_transit_preference,
            "amap_final_transit_vehicle_types": amap_final_transit_vehicle_types,
            "amap_transit_preference_satisfied": amap_transit_preference_satisfied,
            "amap_mode_fallback_used": amap_mode_fallback_used,
            "amap_mode_fallback_reason": amap_mode_fallback_reason,
            "amap_leg_final_modes": amap_leg_final_modes,
            "amap_leg_transit_vehicle_types": amap_leg_transit_vehicle_types,
            "amap_leg_mode_fallbacks": amap_leg_mode_fallbacks,
            "amap_result_reasonable": amap_result_reasonable,
            "amap_latency_ms": amap_latency_ms,
            "amap_distance_meters": amap_distance_meters,
            "amap_duration_seconds": amap_duration_seconds,
            "amap_leg_count": amap_leg_count,
            "amap_available_leg_count": amap_available_leg_count,
            "amap_step_count": amap_step_count,
            "amap_first_steps": amap_first_steps,
            "total_latency_ms": _elapsed_ms(case_started_at),
        }
    )
    score = mean([1.0 if row["intent_correct"] else 0.0, 1.0 if slot_correct else 0.0])
    if amap_checked:
        score = mean([score, 1.0 if amap_result_reasonable else 0.0])
    row["score"] = round(score, 4)
    row["decision"] = "pass" if row["intent_correct"] and slot_correct and (amap_result_reasonable is not False) else "fail"
    return row


async def run_benchmark(
    dataset_path: Path,
    output_cases: Path,
    output_summary: Path,
    *,
    check_amap: bool = False,
    progress_every: int = 0,
    max_cases: int | None = None,
    case_delay_ms: int = 0,
) -> None:
    dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
    cases = list(dataset.get("cases", []))
    if max_cases is not None:
        cases = cases[:max_cases]
    rows = []
    for index, case in enumerate(cases, start=1):
        rows.append(await _resolve_case(case, check_amap=check_amap))
        if progress_every and (index % progress_every == 0 or index == len(cases)):
            print(f"processed {index}/{len(cases)} cases")
        if case_delay_ms > 0 and index < len(cases):
            await asyncio.sleep(case_delay_ms / 1000)
    by_category: dict[str, dict[str, float]] = defaultdict(lambda: {"total": 0.0, "pass": 0.0})
    by_mode: dict[str, dict[str, float]] = defaultdict(lambda: {"total": 0.0, "pass": 0.0})
    by_request_kind: dict[str, dict[str, float]] = defaultdict(lambda: {"total": 0.0, "pass": 0.0})
    for row in rows:
        category = str(row.get("category") or "unknown")
        by_category[category]["total"] += 1
        mode = str(row.get("actual_mode") or row.get("expected_mode") or "unknown")
        by_mode[mode]["total"] += 1
        request_kind = str(row.get("actual_request_kind") or row.get("expected_request_kind") or "unknown")
        by_request_kind[request_kind]["total"] += 1
        if row.get("decision") == "pass":
            by_category[category]["pass"] += 1
            by_mode[mode]["pass"] += 1
            by_request_kind[request_kind]["pass"] += 1
    amap_checked_rows = [row for row in rows if row.get("amap_checked")]
    summary = {
        "dataset_name": dataset.get("dataset_name"),
        "dataset_version": dataset.get("dataset_version"),
        "total_cases": len(rows),
        "pass_count": sum(1 for row in rows if row.get("decision") == "pass"),
        "pass_rate": (sum(1 for row in rows if row.get("decision") == "pass") / len(rows)) if rows else 0,
        "intent_correct_rate": (sum(1 for row in rows if row.get("intent_correct")) / len(rows)) if rows else 0,
        "slot_correct_rate": (
            sum(
                1
                for row in rows
                if row.get("expected_intent") == "navigation_text" and row.get("slot_correct") is True
            ) /
            max(1, sum(1 for row in rows if row.get("expected_intent") == "navigation_text"))
        ),
        "amap_checked_cases": sum(1 for row in rows if row.get("amap_checked")),
        "amap_reasonable_cases": sum(1 for row in rows if row.get("amap_result_reasonable") is True),
        "amap_final_answer_cases": sum(1 for row in rows if row.get("amap_final_answer")),
        "amap_mode_fallback_cases": sum(1 for row in rows if row.get("amap_mode_fallback_used") is True),
        "amap_transit_preference_satisfied_cases": sum(
            1 for row in rows if row.get("amap_transit_preference_satisfied") is True
        ),
        "amap_reasonable_rate": (
            sum(1 for row in amap_checked_rows if row.get("amap_result_reasonable") is True) / len(amap_checked_rows)
            if amap_checked_rows
            else 0
        ),
        "latency_ms": {
            "intent": _latency_stats(rows, "intent_latency_ms"),
            "slot": _latency_stats(rows, "slot_latency_ms"),
            "amap": _latency_stats(rows, "amap_latency_ms"),
            "total": _latency_stats(rows, "total_latency_ms"),
        },
        "by_category": {
            category: {
                "total": int(stats["total"]),
                "pass": int(stats["pass"]),
                "pass_rate": (stats["pass"] / stats["total"]) if stats["total"] else 0,
            }
            for category, stats in by_category.items()
        },
        "by_mode": _pass_rate_summary(by_mode),
        "by_request_kind": _pass_rate_summary(by_request_kind),
    }
    output_cases.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    output_summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


def export_excel(case_path: Path, summary_path: Path, output_path: Path) -> None:
    rows = json.loads(case_path.read_text(encoding="utf-8"))
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_sheet = workbook.create_sheet("summary")
    summary_sheet.append([
        "dataset_name", "dataset_version", "total_cases", "pass_count", "pass_rate",
        "intent_correct_rate", "slot_correct_rate", "amap_checked_cases", "amap_reasonable_cases",
        "amap_final_answer_cases", "amap_mode_fallback_cases",
        "amap_transit_preference_satisfied_cases", "amap_reasonable_rate",
    ])
    summary_sheet.append([
        summary.get("dataset_name"),
        summary.get("dataset_version"),
        summary.get("total_cases"),
        summary.get("pass_count"),
        summary.get("pass_rate"),
        summary.get("intent_correct_rate"),
        summary.get("slot_correct_rate"),
        summary.get("amap_checked_cases"),
        summary.get("amap_reasonable_cases"),
        summary.get("amap_final_answer_cases"),
        summary.get("amap_mode_fallback_cases"),
        summary.get("amap_transit_preference_satisfied_cases"),
        summary.get("amap_reasonable_rate"),
    ])
    latency_sheet = workbook.create_sheet("latency")
    latency_sheet.append(["stage", "count", "avg_ms", "median_ms", "p95_ms", "max_ms"])
    for stage, stats in dict(summary.get("latency_ms") or {}).items():
        latency_sheet.append([
            stage,
            stats.get("count"),
            stats.get("avg"),
            stats.get("median"),
            stats.get("p95"),
            stats.get("max"),
        ])

    case_sheet = workbook.create_sheet("cases")
    headers = [
        "case_id", "category", "sub_category", "query", "language", "mode_case", "distance_bucket",
        "expected_intent", "actual_intent", "intent_correct", "intent_latency_ms",
        "expected_request_kind", "actual_request_kind", "expected_origin", "actual_origin",
        "expected_destinations", "actual_destinations", "expected_source", "actual_source",
        "expected_mode", "actual_mode", "expected_mode_source", "actual_mode_source",
        "expected_transit_preference", "actual_transit_preference",
        "expected_origin_place_id", "actual_origin_place_id", "expected_destination_place_ids",
        "actual_destination_place_ids", "origin_place_source", "destination_place_sources",
        "slot_latency_ms", "slot_debug", "slot_correct",
        "amap_checked", "amap_status", "amap_reason", "amap_final_answer",
        "amap_requested_mode", "amap_final_mode", "amap_requested_transit_preference",
        "amap_final_transit_vehicle_types", "amap_transit_preference_satisfied",
        "amap_mode_fallback_used", "amap_mode_fallback_reason",
        "amap_leg_final_modes", "amap_leg_transit_vehicle_types", "amap_leg_mode_fallbacks",
        "amap_result_reasonable", "amap_latency_ms",
        "amap_distance_meters", "amap_duration_seconds", "amap_leg_count", "amap_available_leg_count",
        "amap_step_count", "amap_first_steps", "total_latency_ms", "score", "decision"
    ]
    case_sheet.append(headers)
    for row in rows:
        serialized = []
        for header in headers:
            value = row.get(header)
            if isinstance(value, (list, dict)):
                serialized.append(json.dumps(value, ensure_ascii=False))
            else:
                serialized.append(value)
        case_sheet.append(serialized)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Run navigation hard-case benchmark")
    parser.add_argument("--dataset", required=True)
    parser.add_argument("--output-cases", required=True)
    parser.add_argument("--output-summary", required=True)
    parser.add_argument("--excel-output")
    parser.add_argument("--check-amap", action="store_true")
    parser.add_argument("--progress-every", type=int, default=0)
    parser.add_argument("--max-cases", type=int)
    parser.add_argument("--case-delay-ms", type=int, default=0)
    args = parser.parse_args()
    asyncio.run(
        run_benchmark(
            Path(args.dataset),
            Path(args.output_cases),
            Path(args.output_summary),
            check_amap=args.check_amap,
            progress_every=args.progress_every,
            max_cases=args.max_cases,
            case_delay_ms=args.case_delay_ms,
        )
    )
    if args.excel_output:
        export_excel(Path(args.output_cases), Path(args.output_summary), Path(args.excel_output))


def _elapsed_ms(started_at: float) -> int:
    return int((perf_counter() - started_at) * 1000)


def _latency_stats(rows: list[dict[str, Any]], key: str) -> dict[str, float | int | None]:
    values = [int(value) for row in rows if (value := row.get(key)) is not None]
    if not values:
        return {"count": 0, "avg": None, "median": None, "p95": None, "max": None}
    sorted_values = sorted(values)
    p95_index = min(len(sorted_values) - 1, int(len(sorted_values) * 0.95))
    return {
        "count": len(sorted_values),
        "avg": round(mean(sorted_values), 2),
        "median": round(median(sorted_values), 2),
        "p95": sorted_values[p95_index],
        "max": max(sorted_values),
    }


def _pass_rate_summary(grouped: dict[str, dict[str, float]]) -> dict[str, dict[str, float | int]]:
    return {
        key: {
            "total": int(stats["total"]),
            "pass": int(stats["pass"]),
            "pass_rate": (stats["pass"] / stats["total"]) if stats["total"] else 0,
        }
        for key, stats in grouped.items()
    }


def _transit_preference_satisfied(
    requested_preference: str | None,
    final_mode: str | None,
    vehicle_types: list[str],
    fallback_used: bool,
) -> bool | None:
    if requested_preference not in {"bus", "subway", "public_transport"}:
        return None
    if fallback_used:
        return True
    if final_mode != "transit":
        return False
    vehicle_type_set = set(vehicle_types)
    if requested_preference == "bus":
        return "bus" in vehicle_type_set and "subway" not in vehicle_type_set
    if requested_preference == "subway":
        return "subway" in vehicle_type_set
    return bool(vehicle_type_set)


if __name__ == "__main__":
    main()

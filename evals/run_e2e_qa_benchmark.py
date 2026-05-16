from __future__ import annotations

import argparse
import asyncio
import json
import time
from pathlib import Path
from statistics import mean
from typing import Any

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

FIRST_STOP_COORDS = {"latitude": 39.9050, "longitude": 116.3976}
LATENCY_STAGE_KEYS = [
    "preprocess_ms",
    "routing_ms",
    "context_build_ms",
    "generation_ms",
    "postprocess_ms",
    "total_ms",
]
CONTEXT_BUILD_BREAKDOWN_KEYS = [
    "attraction_lookup_ms",
    "profile_lookup_ms",
    "live_info_lookup_ms",
    "weather_target_resolution_ms",
    "weather_lookup_ms",
    "navigation_slot_resolution_ms",
    "navigation_lookup_ms",
    "rag_lookup_ms",
    "prompt_projection_ms",
]
INTERMEDIATE_JSON_KEYS = [
    "resolved_query_json",
    "preprocessing_json",
    "rule_result_json",
    "router_fallback_result_json",
    "router_final_result_json",
    "translation_json",
    "answer_translation_json",
    "sql_attraction_json",
    "sql_profile_json",
    "rag_context_json",
    "rag_chunks_json",
    "live_info_json",
    "weather_json",
    "navigation_json",
    "context_build_breakdown_ms",
]


def _lower_list(values: list[str]) -> list[str]:
    return [value.lower() for value in values]


def _limit_cases_per_intent(cases: list[dict[str, Any]], limit: int | None) -> list[dict[str, Any]]:
    if limit is None:
        return cases
    counts: dict[str, int] = {}
    selected: list[dict[str, Any]] = []
    for case in cases:
        intent = str(case.get("expected_intent") or "unknown")
        if counts.get(intent, 0) >= limit:
            continue
        selected.append(case)
        counts[intent] = counts.get(intent, 0) + 1
    return selected


def _match_expected_data_source(expected: str, data: dict[str, Any]) -> bool:
    metadata = data.get("metadata") or {}
    retrieval_strategy = metadata.get("retrieval_strategy")
    degraded_reason = metadata.get("degraded_reason")
    context = metadata.get("context") or {}
    session_context = context.get("session_context") or {}
    if expected == "none":
        return True
    if expected == "sql":
        return retrieval_strategy in {None, "sql_only"} or metadata.get("grounding") == "sql"
    if expected == "sql_then_rag":
        return retrieval_strategy == "sql_then_rag" or metadata.get("rag_query_status") == "ok"
    if expected == "model_knowledge":
        return retrieval_strategy == "model_knowledge" or metadata.get("grounding") == "model_knowledge"
    if expected == "session_context":
        return bool(session_context)
    if expected == "product_redirect":
        return metadata.get("manual_route_edit_redirect") is True
    if expected == "live_search":
        return data.get("intent") == "live_info" and bool(metadata.get("sources"))
    if expected == "live_search_degraded":
        return metadata.get("degraded") is True and degraded_reason == "missing_provider_config"
    if expected == "live_info":
        return data.get("intent") == "live_info" and (
            bool(metadata.get("sources"))
            or metadata.get("status") in {"ok", "degraded", "unavailable"}
            or isinstance(metadata.get("live_info"), dict)
        )
    if expected == "boundary_guard":
        answer = str(data.get("answer") or "").lower()
        boundary_tokens = ["official", "当前不支持", "cannot confirm", "reliable", "官网"]
        return (data.get("supported") is False) or any(token in answer for token in boundary_tokens)
    if expected == "profile_plus_sql":
        return bool(metadata.get("context"))
    if expected == "weather":
        return data.get("intent") == "weather_info" and bool(metadata.get("weather"))
    if expected == "navigation":
        return data.get("intent") == "navigation_text" and bool(metadata.get("navigation"))
    if expected == "translation":
        return data.get("intent") == "translation" and bool(metadata.get("translation"))
    return False


def _score_case(
    case: dict[str, Any],
    payload: dict[str, Any],
    http_status: int,
    api_status: int | None,
) -> tuple[float, str, list[str]]:
    data = payload.get("data") or {}
    answer = str(data.get("answer") or "")
    answer_l = answer.lower()
    score = 0
    notes: list[str] = []

    if http_status == 200 and api_status == 0:
        score += 1
    else:
        notes.append("http_or_api_failure")

    if data.get("intent") == case.get("expected_intent"):
        score += 1
    else:
        notes.append(f"intent_mismatch:{data.get('intent')}")

    if _match_expected_data_source(str(case.get("expected_data_source") or "none"), data):
        score += 1
    else:
        notes.append("data_source_mismatch")

    assertions = case.get("assertions") or {}
    must_include_any = _lower_list(list(assertions.get("must_include_any") or []))
    if not must_include_any or any(token in answer_l for token in must_include_any):
        score += 1
    else:
        notes.append("missing_expected_phrase")

    must_not_depend_on = _lower_list(list(assertions.get("must_not_depend_on") or []))
    if must_not_depend_on and any(token in answer_l for token in must_not_depend_on):
        notes.append("unexpected_dependency_phrase")
    else:
        score += 1

    decision = "pass" if score >= 4 and not notes else "fail"
    return float(score), decision, notes


async def _create_itinerary(
    client: httpx.AsyncClient,
    user_id: str | None,
    title: str,
) -> dict[str, Any]:
    response = await client.post(
        "/api/v1/planning/itineraries",
        json={
            "user_id": user_id,
            "title": title,
            "preferences": {"preferred_poi_count": 2},
        },
    )
    response.raise_for_status()
    return response.json()["data"]


async def _create_guide_session(
    client: httpx.AsyncClient,
    itinerary: dict[str, Any],
    language: str,
    current_stop_index: int | None = None,
) -> str:
    context: dict[str, Any] = {"language": language}
    if current_stop_index is not None:
        context["current_stop_index"] = current_stop_index
    response = await client.post(
        "/api/v1/session/guide",
        json={
            "itinerary_id": itinerary["id"],
            "itinerary_version_id": itinerary["version"]["id"],
            "context": context,
        },
    )
    response.raise_for_status()
    guide_session_id = response.json()["data"]["id"]
    await client.post(f"/api/v1/session/guide/{guide_session_id}/start")
    return guide_session_id


async def _mutate_session_if_needed(
    client: httpx.AsyncClient,
    guide_session_id: str,
    mode: str,
) -> None:
    if mode != "session_mutation":
        return
    await client.post(f"/api/v1/gps/update/{guide_session_id}", json=FIRST_STOP_COORDS)
    await client.post(f"/api/v1/guide/playback/{guide_session_id}", json={"action": "play"})
    await client.post(f"/api/v1/guide/playback/{guide_session_id}", json={"action": "complete"})


async def _prepare_rag(client: httpx.AsyncClient) -> None:
    try:
        await client.post(
            "/api/v1/rag/index-runs/rebuild",
            json={"language": "en", "use_seed": False},
            timeout=120.0,
        )
    except Exception:
        return


def _extract_llm_usage(
    data: dict[str, Any],
) -> tuple[int | None, int | None, int | None, float | None]:
    metadata = data.get("metadata") or {}
    llm = metadata.get("llm") or {}
    usage = llm.get("usage") or {}
    return (
        usage.get("prompt_tokens"),
        usage.get("completion_tokens"),
        usage.get("total_tokens"),
        llm.get("estimated_total_cost") or llm.get("total_cost"),
    )


def _get_nested(payload: dict[str, Any], *keys: str) -> Any:
    current: Any = payload
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def _as_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _latency_fields(metadata: dict[str, Any]) -> dict[str, Any]:
    latency = metadata.get("latency_ms") if isinstance(metadata.get("latency_ms"), dict) else {}
    fields: dict[str, Any] = {}
    for key in LATENCY_STAGE_KEYS:
        fields[key] = latency.get(key)

    context_breakdown = latency.get("context_build_breakdown_ms")
    if not isinstance(context_breakdown, dict):
        context_breakdown = {}
    for key in CONTEXT_BUILD_BREAKDOWN_KEYS:
        fields[key] = context_breakdown.get(key)
    fields["context_build_breakdown_ms"] = context_breakdown
    fields["weather_target_source"] = context_breakdown.get("weather_target_source")
    fields["weather_target_location_name"] = context_breakdown.get("weather_target_location_name")
    fields["navigation_slot_source"] = context_breakdown.get("navigation_slot_source")
    fields["navigation_slot_request_kind"] = context_breakdown.get("navigation_slot_request_kind")
    fields["rag_needed"] = context_breakdown.get("rag_needed")
    return fields


def _dict_child(payload: dict[str, Any], key: str) -> dict[str, Any]:
    value = payload.get(key)
    return value if isinstance(value, dict) else {}


def _router_fields(metadata: dict[str, Any]) -> dict[str, Any]:
    intent_router = _dict_child(metadata, "intent_router")
    rule_result = _dict_child(intent_router, "rule_result")
    final_result = _dict_child(intent_router, "final_result")
    fallback_result = _dict_child(intent_router, "fallback_result")
    boundary_topic = (
        final_result.get("boundary_topic")
        or rule_result.get("boundary_topic")
        or metadata.get("boundary_topic")
    )
    out_of_scope_subtype = (
        final_result.get("out_of_scope_subtype")
        or rule_result.get("out_of_scope_subtype")
        or metadata.get("out_of_scope_subtype")
    )
    return {
        "rule_intent": rule_result.get("intent"),
        "final_intent": final_result.get("intent"),
        "router_confidence": final_result.get("confidence") or rule_result.get("confidence"),
        "router_margin": rule_result.get("margin"),
        "router_needs_fallback": rule_result.get("needs_fallback"),
        "router_fallback_reason": rule_result.get("fallback_reason"),
        "router_fallback_used": bool(fallback_result) and fallback_result.get("intent") is not None,
        "router_fallback_intent": fallback_result.get("intent"),
        "boundary_topic": boundary_topic,
        "out_of_scope_subtype": out_of_scope_subtype,
    }


def _external_latency_fields(metadata: dict[str, Any]) -> dict[str, Any]:
    translation = _dict_child(metadata, "translation")
    answer_translation = _dict_child(metadata, "answer_translation")
    llm = _dict_child(metadata, "llm")
    return {
        "answer_llm_provider": llm.get("provider"),
        "answer_llm_model": llm.get("model"),
        "answer_llm_latency_ms": llm.get("latency_ms"),
        "answer_structured_output_valid": llm.get("structured_output_valid"),
        "translation_latency_ms": translation.get("latency_ms"),
        "translation_execution_path": translation.get("execution_path"),
        "answer_translation_latency_ms": answer_translation.get("latency_ms"),
        "answer_translation_used": answer_translation.get("used_answer_translation"),
        "answer_translation_execution_path": answer_translation.get("execution_path"),
    }


def _intermediate_fields(data: dict[str, Any]) -> dict[str, Any]:
    metadata = data.get("metadata") if isinstance(data.get("metadata"), dict) else {}
    intent_router = _dict_child(metadata, "intent_router")
    rule_result = _dict_child(intent_router, "rule_result")
    fallback_result = _dict_child(intent_router, "fallback_result")
    final_result = _dict_child(intent_router, "final_result")
    preprocessing = _dict_child(intent_router, "preprocessing")
    resolved_query = _dict_child(metadata, "resolved_query")
    translation = _dict_child(metadata, "translation")
    answer_translation = _dict_child(metadata, "answer_translation")
    attraction = _dict_child(metadata, "attraction") or _dict_child(metadata, "retrieval")
    profile = _dict_child(metadata, "profile")
    rag_context = _dict_child(metadata, "rag")
    live_info = {
        "sources": metadata.get("sources"),
        "updated_at": metadata.get("updated_at"),
        "not_confirmed": metadata.get("not_confirmed"),
        "confidence": metadata.get("confidence"),
        "status": metadata.get("status"),
        "reason": metadata.get("reason"),
        "cache_hit": metadata.get("live_info_cache_hit"),
        "cache_key": metadata.get("live_info_cache_key"),
        "cached_at": metadata.get("live_info_cached_at"),
        "cache_ttl_seconds": metadata.get("live_info_cache_ttl_seconds"),
        "info_type": metadata.get("live_info_info_type"),
    }
    weather = _dict_child(metadata, "weather")
    navigation = _dict_child(metadata, "navigation")
    rag_chunks = rag_context.get("chunks") if isinstance(rag_context.get("chunks"), list) else []
    return {
        "resolved_raw_query": resolved_query.get("raw_query"),
        "resolved_effective_query": resolved_query.get("effective_query"),
        "resolved_pivot_query": resolved_query.get("pivot_query"),
        "resolved_user_language": resolved_query.get("user_language"),
        "resolved_processing_language": resolved_query.get("processing_language"),
        "resolved_used_translation_pivot": resolved_query.get("used_translation_pivot"),
        "preprocess_execution_path": preprocessing.get("execution_path"),
        "preprocess_provider": preprocessing.get("provider"),
        "preprocess_model": preprocessing.get("model"),
        "preprocess_degraded": preprocessing.get("degraded"),
        "preprocess_degraded_reason": preprocessing.get("degraded_reason"),
        "translation_status": translation.get("status"),
        "translation_mode": translation.get("mode"),
        "translation_source_language": translation.get("source_language"),
        "translation_target_language": translation.get("target_language"),
        "translation_display_blocks": translation.get("display_blocks"),
        "translation_user_visible_lines": translation.get("user_visible_lines"),
        "answer_translation_target_language": answer_translation.get("target_language"),
        "sql_attraction_name": attraction.get("name"),
        "sql_attraction_category": attraction.get("category"),
        "sql_profile_user_id": profile.get("user_id"),
        "rag_retrieval_mode": rag_context.get("retrieval_mode"),
        "rag_chunk_count": len(rag_chunks),
        "rag_top_chunk_text": _first_rag_chunk_text(rag_chunks),
        "live_info_status": live_info.get("status"),
        "live_info_cache_hit": live_info.get("cache_hit"),
        "live_info_info_type": live_info.get("info_type"),
        "weather_status": weather.get("status"),
        "weather_location_name": weather.get("location_name"),
        "navigation_status": navigation.get("status"),
        "navigation_origin_name": navigation.get("origin_name"),
        "navigation_destination_name": navigation.get("destination_name"),
        "resolved_query_json": resolved_query,
        "preprocessing_json": preprocessing,
        "rule_result_json": rule_result,
        "router_fallback_result_json": fallback_result,
        "router_final_result_json": final_result,
        "translation_json": translation,
        "answer_translation_json": answer_translation,
        "sql_attraction_json": attraction,
        "sql_profile_json": profile,
        "rag_context_json": rag_context,
        "rag_chunks_json": rag_chunks,
        "live_info_json": live_info,
        "weather_json": weather,
        "navigation_json": navigation,
    }


def _first_rag_chunk_text(chunks: list[Any]) -> str | None:
    if not chunks:
        return None
    first = chunks[0]
    if not isinstance(first, dict):
        return None
    text = first.get("text")
    return text if isinstance(text, str) else None


def _numeric_values(rows: list[dict[str, Any]], key: str) -> list[float]:
    return [
        value
        for row in rows
        if (value := _as_number(row.get(key))) is not None
    ]


def _average(rows: list[dict[str, Any]], key: str) -> float | None:
    values = _numeric_values(rows, key)
    return round(mean(values), 2) if values else None


def _percentile(values: list[float], percentile: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    if len(ordered) == 1:
        return round(ordered[0], 2)
    position = (len(ordered) - 1) * percentile
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    fraction = position - lower
    return round(ordered[lower] + (ordered[upper] - ordered[lower]) * fraction, 2)


def _stage_averages(rows: list[dict[str, Any]]) -> dict[str, float | None]:
    keys = [
        "latency_ms",
        *LATENCY_STAGE_KEYS,
        *CONTEXT_BUILD_BREAKDOWN_KEYS,
        "answer_llm_latency_ms",
        "translation_latency_ms",
        "answer_translation_latency_ms",
    ]
    return {key: _average(rows, key) for key in keys}


def _summarize_group(rows: list[dict[str, Any]], group_key: str) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        groups.setdefault(str(row.get(group_key) or "unknown"), []).append(row)
    summaries: list[dict[str, Any]] = []
    for group_value, group_rows in sorted(groups.items()):
        latency_values = _numeric_values(group_rows, "latency_ms")
        pass_count = sum(1 for row in group_rows if row.get("decision") == "pass")
        intent_correct_count = sum(1 for row in group_rows if row.get("intent_correct") is True)
        slowest = max(group_rows, key=lambda row: float(row.get("latency_ms") or 0))
        item = {
            group_key: group_value,
            "total_cases": len(group_rows),
            "pass_count": pass_count,
            "pass_rate": round(pass_count / len(group_rows), 4) if group_rows else 0,
            "intent_correct_count": intent_correct_count,
            "intent_correct_rate": (
                round(intent_correct_count / len(group_rows), 4) if group_rows else 0
            ),
            "p50_latency_ms": _percentile(latency_values, 0.5),
            "p95_latency_ms": _percentile(latency_values, 0.95),
            "slowest_case_id": slowest.get("case_id") if group_rows else None,
        }
        item.update({f"avg_{key}": value for key, value in _stage_averages(group_rows).items()})
        summaries.append(item)
    return summaries


def _build_summary(
    *,
    dataset: dict[str, Any],
    run_id: str,
    model: str,
    provider: str,
    base_url: str,
    rows: list[dict[str, Any]],
) -> dict[str, Any]:
    latency_values = _numeric_values(rows, "latency_ms")
    pass_count = sum(1 for row in rows if row["decision"] == "pass")
    intent_correct_count = sum(1 for row in rows if row.get("intent_correct") is True)
    summary = {
        "run_id": run_id,
        "dataset_name": dataset.get("dataset_name"),
        "dataset_version": dataset.get("dataset_version"),
        "model": model,
        "provider": provider,
        "base_url": base_url,
        "total_cases": len(rows),
        "avg_score": mean([row["score"] for row in rows]) if rows else 0,
        "pass_count": pass_count,
        "pass_rate": (pass_count / len(rows)) if rows else 0,
        "intent_correct_count": intent_correct_count,
        "intent_correct_rate": (intent_correct_count / len(rows)) if rows else 0,
        "avg_latency_ms": _average(rows, "latency_ms"),
        "p50_latency_ms": _percentile(latency_values, 0.5),
        "p95_latency_ms": _percentile(latency_values, 0.95),
        "estimated_total_cost": sum(float(row.get("estimated_total_cost") or 0.0) for row in rows),
        "stage_latency_ms": _stage_averages(rows),
        "by_language": _summarize_group(rows, "language"),
        "by_category": _summarize_group(rows, "category"),
        "by_expected_intent": _summarize_group(rows, "expected_intent"),
        "by_operation": _summarize_group(rows, "operation"),
    }
    return summary


async def run_cases(
    dataset_path: Path,
    base_url: str,
    model_label: str,
    run_id: str,
    output_cases: Path,
    output_summary: Path,
    llm_provider: str | None = None,
    llm_model: str | None = None,
    intent_filter: set[str] | None = None,
    max_cases: int | None = None,
    max_cases_per_intent: int | None = None,
    progress_every: int = 0,
    case_delay_ms: int = 0,
) -> None:
    dataset = json.loads(dataset_path.read_text(encoding="utf-8"))
    cases = list(dataset["cases"])
    if intent_filter:
        cases = [case for case in cases if case.get("expected_intent") in intent_filter]
    cases = _limit_cases_per_intent(cases, max_cases_per_intent)
    if max_cases is not None:
        cases = cases[:max_cases]
    rows: list[dict[str, Any]] = []
    async with httpx.AsyncClient(base_url=base_url, timeout=90.0) as client:
        if any(case.get("expected_data_source") == "sql_then_rag" for case in cases):
            await _prepare_rag(client)
        for index, case in enumerate(cases, 1):
            request = case.get("request") or case.get("request_template") or {}
            profile_id = case.get("profile_id")
            guide_session_id: str | None = None
            session_state = "none"
            if case.get("session_required"):
                itinerary = await _create_itinerary(
                    client,
                    profile_id,
                    case.get("title") or case.get("case_id"),
                )
                guide_session_id = await _create_guide_session(
                    client,
                    itinerary,
                    request.get("language", case.get("language", "en")),
                    current_stop_index=0,
                )
                await _mutate_session_if_needed(
                    client,
                    guide_session_id,
                    case.get("mode", "single_turn"),
                )
                session_state = case.get("mode", "session_aware")

            qa_payload = {
                "guide_session_id": guide_session_id,
                "user_id": profile_id,
                "query": request.get("query"),
                "language": request.get("language", case.get("language", "en")),
                "llm_provider": llm_provider,
                "llm_model": llm_model,
                "context": request.get("context") if isinstance(request.get("context"), dict) else {},
            }
            started = time.perf_counter()
            response = await client.post("/api/v1/qa/ask", json=qa_payload)
            latency_ms = (time.perf_counter() - started) * 1000
            payload = response.json()
            data = payload.get("data") or {}
            prompt_tokens, completion_tokens, total_tokens, total_cost = _extract_llm_usage(data)
            score, decision, notes = _score_case(
                case,
                payload,
                response.status_code,
                payload.get("code"),
            )
            intent_correct = data.get("intent") == case.get("expected_intent")
            metadata = data.get("metadata") or {}
            latency_fields = _latency_fields(metadata)
            router_fields = _router_fields(metadata)
            external_latency_fields = _external_latency_fields(metadata)
            intermediate_fields = _intermediate_fields(data)
            row = {
                "run_id": run_id,
                "case_id": case.get("case_id"),
                "query_id": case.get("case_id"),
                "category": case.get("category"),
                "language": case.get("language"),
                "request_language": request.get("language", case.get("language", "en")),
                "title": case.get("title"),
                "mode": case.get("mode"),
                "operation": case.get("operation"),
                "query_zh": case.get("query_zh"),
                "expected_intent": case.get("expected_intent"),
                "expected_data_source": case.get("expected_data_source"),
                "session_required": case.get("session_required"),
                "profile_id": profile_id,
                "fixture_mode": case.get("fixture_mode"),
                "variant_group": case.get("variant_group"),
                "session_state": session_state,
                "provider": llm_provider or "api:/api/v1/qa/ask",
                "model": llm_model or model_label,
                "original_query": request.get("query"),
                "request_text": request.get("query"),
                "response_text": data.get("answer"),
                "http_status": response.status_code,
                "api_status": payload.get("code"),
                "latency_ms": round(latency_ms, 2),
                "reported_total_ms": latency_fields.get("total_ms"),
                "preprocess_ms": latency_fields.get("preprocess_ms"),
                "routing_ms": latency_fields.get("routing_ms"),
                "context_build_ms": latency_fields.get("context_build_ms"),
                "generation_ms": latency_fields.get("generation_ms"),
                "postprocess_ms": latency_fields.get("postprocess_ms"),
                "attraction_lookup_ms": latency_fields.get("attraction_lookup_ms"),
                "profile_lookup_ms": latency_fields.get("profile_lookup_ms"),
                "live_info_lookup_ms": latency_fields.get("live_info_lookup_ms"),
                "weather_target_resolution_ms": latency_fields.get("weather_target_resolution_ms"),
                "weather_lookup_ms": latency_fields.get("weather_lookup_ms"),
                "weather_target_source": latency_fields.get("weather_target_source"),
                "weather_target_location_name": latency_fields.get("weather_target_location_name"),
                "navigation_slot_resolution_ms": latency_fields.get(
                    "navigation_slot_resolution_ms"
                ),
                "navigation_lookup_ms": latency_fields.get("navigation_lookup_ms"),
                "navigation_slot_source": latency_fields.get("navigation_slot_source"),
                "navigation_slot_request_kind": latency_fields.get("navigation_slot_request_kind"),
                "rag_needed": latency_fields.get("rag_needed"),
                "rag_lookup_ms": latency_fields.get("rag_lookup_ms"),
                "prompt_projection_ms": latency_fields.get("prompt_projection_ms"),
                "context_build_breakdown_ms": latency_fields.get("context_build_breakdown_ms"),
                **router_fields,
                **external_latency_fields,
                **intermediate_fields,
                "prompt_tokens": prompt_tokens,
                "completion_tokens": completion_tokens,
                "total_tokens": total_tokens,
                "estimated_total_cost": total_cost,
                "intent_correct": intent_correct,
                "score": score,
                "max_score": 5,
                "decision": decision,
                "confidence": metadata.get("confidence"),
                "degraded": metadata.get("degraded"),
                "degraded_reason": metadata.get("degraded_reason") or metadata.get("reason"),
                "retrieval_strategy": metadata.get("retrieval_strategy"),
                "grounding": metadata.get("grounding"),
                "rag_backend_ready": metadata.get("rag_backend_ready"),
                "rag_query_status": metadata.get("rag_query_status"),
                "supported": data.get("supported"),
                "artifact_path": str(output_cases),
                "notes": "; ".join(notes),
            }
            rows.append(row)
            if progress_every > 0 and index % progress_every == 0:
                print(f"processed {index}/{len(cases)} cases")
            if case_delay_ms > 0:
                await asyncio.sleep(case_delay_ms / 1000)

    output_cases.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    summary = _build_summary(
        dataset=dataset,
        run_id=run_id,
        model=llm_model or model_label,
        provider=llm_provider or "api:/api/v1/qa/ask",
        base_url=base_url,
        rows=rows,
    )
    output_summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


def export_excel(case_files: list[Path], summary_files: list[Path], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_sheet = workbook.create_sheet("summary")
    summary_headers = [
        "run_id",
        "dataset_name",
        "dataset_version",
        "provider",
        "model",
        "base_url",
        "total_cases",
        "avg_score",
        "pass_count",
        "pass_rate",
        "intent_correct_count",
        "intent_correct_rate",
        "avg_latency_ms",
        "p50_latency_ms",
        "p95_latency_ms",
        "estimated_total_cost",
    ]
    summary_sheet.append(summary_headers)
    for path in summary_files:
        item = json.loads(path.read_text(encoding="utf-8"))
        summary_sheet.append([_excel_value(item.get(header)) for header in summary_headers])

    case_sheet = workbook.create_sheet("cases")
    case_headers = [
        "run_id",
        "case_id",
        "category",
        "language",
        "request_language",
        "title",
        "mode",
        "operation",
        "original_query",
        "query_zh",
        "expected_intent",
        "expected_data_source",
        "profile_id",
        "fixture_mode",
        "provider",
        "model",
        "request_text",
        "response_text",
        "resolved_raw_query",
        "resolved_effective_query",
        "resolved_pivot_query",
        "resolved_user_language",
        "resolved_processing_language",
        "resolved_used_translation_pivot",
        "http_status",
        "api_status",
        "latency_ms",
        "reported_total_ms",
        "preprocess_ms",
        "routing_ms",
        "context_build_ms",
        "generation_ms",
        "postprocess_ms",
        "attraction_lookup_ms",
        "profile_lookup_ms",
        "live_info_lookup_ms",
        "weather_target_resolution_ms",
        "weather_lookup_ms",
        "weather_target_source",
        "navigation_slot_resolution_ms",
        "navigation_lookup_ms",
        "navigation_slot_source",
        "navigation_slot_request_kind",
        "rag_needed",
        "rag_lookup_ms",
        "prompt_projection_ms",
        "router_needs_fallback",
        "router_fallback_used",
        "router_fallback_reason",
        "rule_intent",
        "final_intent",
        "answer_llm_provider",
        "answer_llm_model",
        "answer_llm_latency_ms",
        "translation_latency_ms",
        "translation_execution_path",
        "answer_translation_latency_ms",
        "answer_translation_used",
        "answer_translation_execution_path",
        "preprocess_execution_path",
        "preprocess_provider",
        "preprocess_model",
        "preprocess_degraded",
        "preprocess_degraded_reason",
        "translation_status",
        "translation_mode",
        "translation_source_language",
        "translation_target_language",
        "translation_display_blocks",
        "translation_user_visible_lines",
        "answer_translation_target_language",
        "sql_attraction_name",
        "sql_attraction_category",
        "sql_profile_user_id",
        "rag_retrieval_mode",
        "rag_chunk_count",
        "rag_top_chunk_text",
        "live_info_status",
        "live_info_cache_hit",
        "live_info_info_type",
        "weather_status",
        "weather_location_name",
        "navigation_status",
        "navigation_origin_name",
        "navigation_destination_name",
        *INTERMEDIATE_JSON_KEYS,
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
        "estimated_total_cost",
        "intent_correct",
        "score",
        "max_score",
        "decision",
        "degraded",
        "degraded_reason",
        "retrieval_strategy",
        "grounding",
        "rag_backend_ready",
        "rag_query_status",
        "supported",
        "notes",
    ]
    case_sheet.append(case_headers)
    for path in case_files:
        rows = json.loads(path.read_text(encoding="utf-8"))
        for row in rows:
            merged = {
                **row,
                "original_query": row.get("original_query") or row.get("request_text"),
            }
            case_sheet.append([_excel_value(merged.get(header)) for header in case_headers])

    review_sheet = workbook.create_sheet("review_cases")
    review_headers = [
        "case_id",
        "语种",
        "原始 query",
        "中文释义",
        "最终 answer",
        "总时延(ms)",
        "结论",
        "评分",
        "意图(预期->实际)",
        "模型",
        "provider",
        "链路",
        "grounding",
        "预处理结果",
        "预处理(ms)",
        "路由结果",
        "路由(ms)",
        "SQL结果",
        "SQL(ms)",
        "RAG结果",
        "RAG(ms)",
        "实时信息结果",
        "实时信息(ms)",
        "天气结果",
        "天气(ms)",
        "导航结果",
        "导航(ms)",
        "生成结果",
        "生成(ms)",
        "翻译结果",
        "翻译(ms)",
        "回译结果",
        "回译(ms)",
        "备注",
    ]
    review_sheet.append(review_headers)
    for path in case_files:
        rows = json.loads(path.read_text(encoding="utf-8"))
        for row in rows:
            merged = {
                **row,
                "original_query": row.get("original_query") or row.get("request_text"),
            }
            human_row = _build_human_review_row(merged)
            review_sheet.append([_excel_value(human_row.get(header)) for header in review_headers])
    _apply_human_review_layout(review_sheet)

    _write_group_sheet(workbook, "latency_by_language", summary_files, "by_language")
    _write_group_sheet(workbook, "latency_by_category", summary_files, "by_category")
    _write_group_sheet(workbook, "latency_by_intent", summary_files, "by_expected_intent")
    _write_group_sheet(workbook, "latency_by_operation", summary_files, "by_operation")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _write_group_sheet(
    workbook: Workbook,
    sheet_name: str,
    summary_files: list[Path],
    group_key: str,
) -> None:
    sheet = workbook.create_sheet(sheet_name)
    group_column = {
        "by_language": "language",
        "by_category": "category",
        "by_expected_intent": "expected_intent",
        "by_operation": "operation",
    }[group_key]
    headers = [
        "run_id",
        "provider",
        "model",
        group_column,
        "total_cases",
        "pass_rate",
        "intent_correct_rate",
        "avg_latency_ms",
        "p50_latency_ms",
        "p95_latency_ms",
        "avg_preprocess_ms",
        "avg_routing_ms",
        "avg_context_build_ms",
        "avg_generation_ms",
        "avg_postprocess_ms",
        "avg_attraction_lookup_ms",
        "avg_profile_lookup_ms",
        "avg_live_info_lookup_ms",
        "avg_weather_target_resolution_ms",
        "avg_weather_lookup_ms",
        "avg_navigation_slot_resolution_ms",
        "avg_navigation_lookup_ms",
        "avg_rag_lookup_ms",
        "avg_answer_llm_latency_ms",
        "avg_translation_latency_ms",
        "avg_answer_translation_latency_ms",
        "slowest_case_id",
    ]
    sheet.append(headers)
    for path in summary_files:
        summary = json.loads(path.read_text(encoding="utf-8"))
        for row in summary.get(group_key) or []:
            merged = {
                "run_id": summary.get("run_id"),
                "provider": summary.get("provider"),
                "model": summary.get("model"),
                **row,
            }
            sheet.append([_excel_value(merged.get(header)) for header in headers])


def _excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _jsonish(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    stripped = value.strip()
    if not stripped or stripped[0] not in "[{":
        return value
    try:
        return json.loads(stripped)
    except json.JSONDecodeError:
        return value


def _readable_text(value: Any, limit: int = 320) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        text = json.dumps(value, ensure_ascii=False)
    else:
        text = str(value)
    text = " ".join(text.split())
    if limit > 0 and len(text) > limit:
        return f"{text[: limit - 3]}..."
    return text


def _ms_sum(*values: Any) -> float | None:
    numbers: list[float] = []
    for value in values:
        if isinstance(value, int | float):
            numbers.append(float(value))
    if not numbers:
        return None
    return round(sum(numbers), 2)


def _summarize_preprocess(row: dict[str, Any]) -> str:
    raw = row.get("resolved_raw_query") or row.get("original_query") or row.get("request_text")
    effective = row.get("resolved_effective_query")
    pivot = row.get("resolved_pivot_query")
    path = row.get("preprocess_execution_path")
    processing_language = row.get("resolved_processing_language")
    if row.get("resolved_used_translation_pivot"):
        return _readable_text(
            f"pivot used; raw={raw}; pivot={pivot}; processing={processing_language}; path={path}"
        )
    return _readable_text(
        f"no pivot; effective={effective or raw}; processing={processing_language}; path={path}"
    )


def _summarize_routing(row: dict[str, Any]) -> str:
    parts = [
        f"expected={row.get('expected_intent')}",
        f"rule={row.get('rule_intent')}",
        f"final={row.get('final_intent')}",
        f"fallback={bool(row.get('router_fallback_used'))}",
    ]
    if row.get("router_confidence") is not None:
        parts.append(f"confidence={row.get('router_confidence')}")
    if row.get("router_fallback_reason"):
        parts.append(f"reason={row.get('router_fallback_reason')}")
    return "; ".join(parts)


def _summarize_sql(row: dict[str, Any]) -> str:
    attraction = row.get("sql_attraction_name")
    category = row.get("sql_attraction_category")
    profile = row.get("sql_profile_user_id")
    sql_json = _jsonish(row.get("sql_attraction_json"))
    source = sql_json.get("source") if isinstance(sql_json, dict) else None
    parts = []
    if attraction:
        parts.append(f"attraction={attraction}")
    if category:
        parts.append(f"category={category}")
    if profile:
        parts.append(f"profile={profile}")
    if source:
        parts.append(f"source={source}")
    return "; ".join(parts) if parts else "not used"


def _summarize_rag(row: dict[str, Any]) -> str:
    mode = row.get("rag_retrieval_mode") or row.get("retrieval_strategy")
    chunks = row.get("rag_chunk_count")
    top = row.get("rag_top_chunk_text")
    if chunks:
        return _readable_text(f"mode={mode}; chunks={chunks}; top={top}")
    if mode:
        return f"mode={mode}; chunks=0"
    return "not used"


def _summarize_live_info(row: dict[str, Any]) -> str:
    live = _jsonish(row.get("live_info_json"))
    if not isinstance(live, dict) or not live:
        return "not used"
    sources = live.get("sources") or []
    domains = []
    for source in sources[:3]:
        if isinstance(source, dict) and source.get("domain"):
            domains.append(str(source["domain"]))
    reason = live.get("reason")
    return _readable_text(
        f"status={live.get('status')}; confidence={live.get('confidence')}; "
        f"cache_hit={live.get('cache_hit')}; info_type={live.get('info_type')}; "
        f"sources={len(sources)} {','.join(domains)}; reason={reason}"
    )


def _summarize_weather(row: dict[str, Any]) -> str:
    weather = _jsonish(row.get("weather_json"))
    if not isinstance(weather, dict) or not weather:
        return "not used"
    return _readable_text(
        f"status={weather.get('status')}; location={weather.get('location_name')}; "
        f"weather={weather.get('weather')}; temp={weather.get('temperature_celsius')}; "
        f"humidity={weather.get('humidity')}; report_time={weather.get('report_time')}"
    )


def _summarize_navigation(row: dict[str, Any]) -> str:
    navigation = _jsonish(row.get("navigation_json"))
    if not isinstance(navigation, dict) or not navigation:
        return "not used"
    return _readable_text(
        f"status={navigation.get('status')}; {navigation.get('origin_name')} -> "
        f"{navigation.get('destination_name')}; distance_m={navigation.get('distance_meters')}; "
        f"duration_s={navigation.get('duration_seconds')}"
    )


def _summarize_generation(row: dict[str, Any]) -> str:
    provider = row.get("answer_llm_provider") or row.get("provider")
    model = row.get("answer_llm_model") or row.get("model")
    structured = row.get("answer_structured_output_valid")
    return f"provider={provider}; model={model}; structured_valid={structured}"


def _summarize_translation(row: dict[str, Any]) -> str:
    translation = _jsonish(row.get("translation_json"))
    if isinstance(translation, dict) and translation:
        lines = translation.get("user_visible_lines") or row.get("translation_user_visible_lines")
        return _readable_text(
            f"status={translation.get('status')}; mode={translation.get('mode')}; "
            f"{translation.get('source_language')}->{translation.get('target_language')}; "
            f"lines={lines}"
        )
    if row.get("translation_status"):
        return _readable_text(
            f"status={row.get('translation_status')}; mode={row.get('translation_mode')}; "
            f"{row.get('translation_source_language')}->{row.get('translation_target_language')}"
        )
    return "not used"


def _summarize_answer_translation(row: dict[str, Any]) -> str:
    answer_translation = _jsonish(row.get("answer_translation_json"))
    if isinstance(answer_translation, dict) and answer_translation:
        source_language = answer_translation.get("source_language")
        target_language = answer_translation.get("target_language")
        execution_path = answer_translation.get("execution_path")
        degraded = answer_translation.get("degraded")
        return _readable_text(
            f"used={answer_translation.get('used_answer_translation')}; "
            f"{source_language}->{target_language}; path={execution_path}; degraded={degraded}"
        )
    return (
        f"used={bool(row.get('answer_translation_used'))}; "
        f"path={row.get('answer_translation_execution_path')}"
    )


def _build_human_review_row(row: dict[str, Any]) -> dict[str, Any]:
    original_query = row.get("original_query") or row.get("request_text")
    total_latency = row.get("reported_total_ms") or row.get("latency_ms")
    intent_match = f"{row.get('expected_intent')} -> {row.get('final_intent') or row.get('intent')}"
    return {
        "case_id": row.get("case_id"),
        "语种": row.get("language"),
        "原始 query": original_query,
        "中文释义": row.get("query_zh"),
        "最终 answer": row.get("response_text"),
        "总时延(ms)": total_latency,
        "结论": row.get("decision"),
        "评分": f"{row.get('score')}/{row.get('max_score')}",
        "意图(预期->实际)": intent_match,
        "模型": row.get("model"),
        "provider": row.get("provider"),
        "链路": row.get("operation"),
        "grounding": row.get("grounding") or row.get("retrieval_strategy"),
        "预处理结果": _summarize_preprocess(row),
        "预处理(ms)": row.get("preprocess_ms"),
        "路由结果": _summarize_routing(row),
        "路由(ms)": row.get("routing_ms"),
        "SQL结果": _summarize_sql(row),
        "SQL(ms)": _ms_sum(row.get("attraction_lookup_ms"), row.get("profile_lookup_ms")),
        "RAG结果": _summarize_rag(row),
        "RAG(ms)": row.get("rag_lookup_ms"),
        "实时信息结果": _summarize_live_info(row),
        "实时信息(ms)": row.get("live_info_lookup_ms"),
        "天气结果": _summarize_weather(row),
        "天气(ms)": row.get("weather_lookup_ms"),
        "导航结果": _summarize_navigation(row),
        "导航(ms)": row.get("navigation_lookup_ms"),
        "生成结果": _summarize_generation(row),
        "生成(ms)": row.get("generation_ms"),
        "翻译结果": _summarize_translation(row),
        "翻译(ms)": row.get("translation_latency_ms"),
        "回译结果": _summarize_answer_translation(row),
        "回译(ms)": row.get("answer_translation_latency_ms"),
        "备注": row.get("notes"),
    }


def _apply_human_review_layout(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "A": 24,
        "B": 8,
        "C": 42,
        "D": 36,
        "E": 80,
        "F": 14,
        "I": 24,
        "J": 30,
        "K": 18,
        "L": 32,
        "M": 48,
        "O": 48,
        "Q": 42,
        "S": 50,
        "U": 50,
        "W": 50,
        "Y": 50,
        "AA": 42,
        "AC": 50,
        "AE": 42,
        "AG": 36,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    parser = argparse.ArgumentParser(description="Run E2E QA benchmark against a single base URL")
    parser.add_argument("--dataset")
    parser.add_argument("--base-url")
    parser.add_argument("--model-label")
    parser.add_argument("--llm-provider")
    parser.add_argument("--llm-model")
    parser.add_argument("--run-id")
    parser.add_argument("--output-cases")
    parser.add_argument("--output-summary")
    parser.add_argument("--sweep")
    parser.add_argument("--output-dir")
    parser.add_argument("--export-only", action="store_true")
    parser.add_argument("--case-files", nargs="*")
    parser.add_argument("--summary-files", nargs="*")
    parser.add_argument("--excel-output")
    parser.add_argument("--intent-filter", nargs="*")
    parser.add_argument("--max-cases", type=int)
    parser.add_argument("--max-cases-per-intent", type=int)
    parser.add_argument("--progress-every", type=int, default=0)
    parser.add_argument("--case-delay-ms", type=int, default=0)
    args = parser.parse_args()

    if args.export_only:
        export_excel(
            [Path(p) for p in args.case_files or []],
            [Path(p) for p in args.summary_files or []],
            Path(args.excel_output),
        )
        return

    if args.sweep:
        sweep = json.loads(Path(args.sweep).read_text(encoding="utf-8"))
        output_dir = Path(args.output_dir or "evals/results/e2e_qa_model_sweep")
        output_dir.mkdir(parents=True, exist_ok=True)
        case_files: list[Path] = []
        summary_files: list[Path] = []
        for item in sweep.get("models", []):
            label = str(item["label"])
            provider = str(item["provider"])
            model = str(item["model"])
            case_path = output_dir / f"{label}_cases.json"
            summary_path = output_dir / f"{label}_summary.json"
            asyncio.run(
                run_cases(
                    dataset_path=Path(args.dataset),
                    base_url=args.base_url,
                    model_label=label,
                    run_id=f"{args.run_id}-{label}",
                    output_cases=case_path,
                    output_summary=summary_path,
                    llm_provider=provider,
                    llm_model=model,
                    intent_filter=set(args.intent_filter or []) or None,
                    max_cases=args.max_cases,
                    max_cases_per_intent=args.max_cases_per_intent,
                    progress_every=args.progress_every,
                    case_delay_ms=args.case_delay_ms,
                )
            )
            case_files.append(case_path)
            summary_files.append(summary_path)
        export_excel(case_files, summary_files, output_dir / "e2e_model_sweep_results.xlsx")
        return

    required_args = [
        args.dataset,
        args.base_url,
        args.model_label,
        args.run_id,
        args.output_cases,
        args.output_summary,
    ]
    if not all(required_args):
        raise SystemExit(
            "dataset/base-url/model-label/run-id/output-cases/output-summary are required "
            "unless --export-only or --sweep is used"
        )

    asyncio.run(
        run_cases(
            dataset_path=Path(args.dataset),
            base_url=args.base_url,
            model_label=args.model_label,
            run_id=args.run_id,
            output_cases=Path(args.output_cases),
            output_summary=Path(args.output_summary),
            llm_provider=args.llm_provider,
            llm_model=args.llm_model,
            intent_filter=set(args.intent_filter or []) or None,
            max_cases=args.max_cases,
            max_cases_per_intent=args.max_cases_per_intent,
            progress_every=args.progress_every,
            case_delay_ms=args.case_delay_ms,
        )
    )


if __name__ == "__main__":
    main()

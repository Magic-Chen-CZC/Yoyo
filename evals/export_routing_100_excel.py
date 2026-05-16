from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "results"
OUTPUT = RESULTS / "routing_benchmark_100_results.xlsx"

CASE_FILES = [
    RESULTS / "routing_rules_only_100_cases.json",
    RESULTS / "routing_rules_plus_fallback_gpt54nano_100_cases.json",
]
SUMMARY_FILES = [
    RESULTS / "routing_rules_only_100_summary.json",
    RESULTS / "routing_rules_plus_fallback_gpt54nano_100_summary.json",
]

ROUTING_CASE_HEADERS = [
    "candidate_name",
    "case_id",
    "bucket",
    "language",
    "query",
    "query_zh",
    "expected_intent",
    "rule_intent",
    "final_intent",
    "intent_correct",
    "domain_supported",
    "hard_deny",
    "needs_fallback",
    "fallback_used",
    "fallback_reason",
    "boundary_topic",
    "out_of_scope_subtype",
    "latency_ms",
    "prompt_tokens",
    "completion_tokens",
    "total_tokens",
    "estimated_total_cost",
    "fallback_provider",
    "fallback_model",
    "notes",
]

ROUTING_SUMMARY_HEADERS = [
    "dataset_name",
    "dataset_version",
    "candidate_name",
    "fallback_enabled",
    "total_cases",
    "correct_cases",
    "accuracy",
    "fallback_needed_cases",
    "fallback_used_cases",
    "fallback_hit_rate",
    "needs_fallback_accuracy",
    "boundary_accuracy",
    "manual_route_edit_accuracy",
    "conflict_accuracy",
    "avg_latency_ms",
    "estimated_total_cost",
]


def _load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


workbook = Workbook()
workbook.remove(workbook.active)

summary_sheet = workbook.create_sheet("routing_summary")
summary_sheet.append(ROUTING_SUMMARY_HEADERS)
for path in SUMMARY_FILES:
    item = _load_json(path)
    summary_sheet.append([item.get(header) for header in ROUTING_SUMMARY_HEADERS])

case_sheet = workbook.create_sheet("routing_cases")
case_sheet.append(ROUTING_CASE_HEADERS)
for path in CASE_FILES:
    rows = _load_json(path)
    for item in rows:
        case_sheet.append([item.get(header) for header in ROUTING_CASE_HEADERS])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
workbook.save(OUTPUT)
print(str(OUTPUT))
